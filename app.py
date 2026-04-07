from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from werkzeug.utils import secure_filename
import pandas as pd
import psycopg2
from psycopg2.extensions import ISOLATION_LEVEL_AUTOCOMMIT
from haystack.utils import Secret
from haystack_integrations.components.generators.google_ai import GoogleAIGeminiGenerator
from io import BytesIO
import pdfplumber
import PyPDF2
import re
import os
import json
import logging
from datetime import datetime
import unicodedata
import numpy as np
from flask.json.provider import DefaultJSONProvider

# Custom JSON provider to handle NumPy types
class CustomJSONProvider(DefaultJSONProvider):
    def default(self, o):
        if isinstance(o, np.integer):
            return int(o)
        if isinstance(o, np.floating):
            return float(o)
        if isinstance(o, np.ndarray):
            return o.tolist()
        return super().default(o)

# Configure logging
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

# Initialize Flask with custom JSON provider
app = Flask(__name__, static_folder=os.path.abspath(os.path.join(os.path.dirname(__file__), 'static')), static_url_path='')
app.json = CustomJSONProvider(app)
CORS(app)

# Load environment variables
from dotenv import load_dotenv
load_dotenv()

try:
    api_key = os.getenv('GOOGLE_AI_API_KEY')
    if not api_key:
        raise ValueError("GOOGLE_AI_API_KEY environment variable is required")
    
    gemini = GoogleAIGeminiGenerator(
        model="gemini-2.5-flash", 
        api_key=Secret.from_token(api_key)
    )
    logger.debug("Gemini Generator initialized successfully")
except Exception as e:
    logger.error(f"Failed to initialize Gemini Generator: {str(e)}")
    raise Exception(f"Failed to initialize Gemini Generator: {str(e)}")

# Store multiple tables and current table
table_metadata = {
    "current_table": None,
    "tables": {}  # {table_name: column_metadata}
}

def sanitize_table_name(name: str) -> str:
    sanitized_name = re.sub(r'[^a-zA-Z0-9_]', '_', name)
    sanitized_name = sanitized_name.lower()
    sanitized_name = sanitized_name.strip('_')
    return sanitized_name

def connect_to_postgres(dbname=None):
    try:
        # Neon DB connection string
        neon_conn_str = "postgresql://neondb_owner:npg_5Y7zmyQiaMSZ@ep-nameless-paper-a4x2j73a-pooler.us-east-1.aws.neon.tech/neondb?sslmode=require&channel_binding=require"
        connection = psycopg2.connect(neon_conn_str)
        connection.set_isolation_level(ISOLATION_LEVEL_AUTOCOMMIT)
        logger.debug("Connected to Neon PostgreSQL successfully")
        return connection
    except Exception as e:
        logger.error(f"Error connecting to Neon PostgreSQL: {str(e)}")
        return {"error": f"Error connecting to Neon PostgreSQL: {str(e)}"}

def setup_query_history_table():
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        logger.error(f"Failed to connect to PostgreSQL: {connection['error']}")
        return False
    
    cursor = connection.cursor()
    
    # Create table if it doesn't exist
    create_table_query = """
    CREATE TABLE IF NOT EXISTS query_history (
        id SERIAL PRIMARY KEY,
        table_name TEXT,
        question TEXT,
        sql_query TEXT,
        results JSONB,
        timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """
    
    # Add missing feedback column for simple like/dislike
    add_simple_feedback = """
    DO $$ 
    BEGIN
        IF NOT EXISTS (SELECT 1 FROM information_schema.columns 
                      WHERE table_name='query_history' AND column_name='feedback') THEN
            ALTER TABLE query_history ADD COLUMN feedback INTEGER DEFAULT NULL;
        END IF;
        
        IF NOT EXISTS (SELECT 1 FROM information_schema.columns 
                      WHERE table_name='query_history' AND column_name='feedback_timestamp') THEN
            ALTER TABLE query_history ADD COLUMN feedback_timestamp TIMESTAMP DEFAULT NULL;
        END IF;
    END $$;
    """
    
    # Create feedback patterns table for RLHF
    create_feedback_table = """
    CREATE TABLE IF NOT EXISTS feedback_patterns (
        id SERIAL PRIMARY KEY,
        question_pattern TEXT UNIQUE,
        successful_query TEXT,
        feedback_score FLOAT DEFAULT 0.0,
        usage_count INTEGER DEFAULT 1,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """
    
    try:
        cursor.execute(create_table_query)
        cursor.execute(add_simple_feedback)
        
        # Enhanced feedback columns
        add_enhanced_feedback = """
        DO $$ 
        BEGIN
            IF NOT EXISTS (SELECT 1 FROM information_schema.columns 
                          WHERE table_name='query_history' AND column_name='rating') THEN
                ALTER TABLE query_history ADD COLUMN rating INTEGER DEFAULT NULL;
            END IF;
            
            IF NOT EXISTS (SELECT 1 FROM information_schema.columns 
                          WHERE table_name='query_history' AND column_name='categories') THEN
                ALTER TABLE query_history ADD COLUMN categories TEXT[] DEFAULT NULL;
            END IF;
            
            IF NOT EXISTS (SELECT 1 FROM information_schema.columns 
                          WHERE table_name='query_history' AND column_name='comment') THEN
                ALTER TABLE query_history ADD COLUMN comment TEXT DEFAULT NULL;
            END IF;
        END $$;
        """
        cursor.execute(add_enhanced_feedback)
        cursor.execute(create_feedback_table)
        connection.commit()
        logger.debug("Query history and feedback tables created/updated successfully")
        return True
    except Exception as e:
        connection.rollback()
        logger.error(f"Error creating/updating tables: {str(e)}")
        return False
    finally:
        cursor.close()
        connection.close()

# Initialize query history table
setup_query_history_table()

def extract_features(text: str) -> list:
    """Extract resume-specific features (skills, experience, etc.)."""
    text = unicodedata.normalize('NFKD', text.strip().lower())  # Normalize Unicode
    text = re.sub(r'\s+', ' ', text)  # Clean whitespace
    feature_patterns = [
        r'•\s+(.+?)(?=\n|$|•)',  # Bullet points
        r'-\s+(.+?)(?=\n|$|-)',  # Hyphen lists
        r'\*\s+(.+?)(?=\n|$|\*)',  # Asterisk lists
        r'\d+\.\s+(.+?)(?=\n|$|\d+\.)',  # Numbered lists
        r'(?:skill[s]?|proficienc(?:y|ies)|expertise|qualification[s]?|experience|education):?\s*(.+?)(?=\n\n|$)',  # Resume labels
        r'(?:include[s]?|list|provid(?:e|es)|offer[s]?):?\s+(.+?)(?=\n\n|$)',  # Common verbs
        r'(?:key|core|primary)\s+(?:skill[s]?|qualification[s]?):?\s+(.+?)(?=\n\n|$)',  # Key skills
        r'^\s*[^:\n]+:\s+(.+?)(?=\n\n|$)',  # Colon-separated lists
        r'[^.!?]*\b(skill|proficiency|expertise|qualification|experience|education)\b[^.!?]*?(?=\n|$)',  # Keyword sentences
        r'([^,;\n|]+?)(?:\s*[|,;]\s*([^,;\n|]+?))+(?=\n|$)'  # Comma/pipe-separated lists
    ]
    features = []
    for pattern in feature_patterns:
        matches = re.findall(pattern, text, re.MULTILINE | re.IGNORECASE)
        for match in matches:
            if isinstance(match, tuple):
                match = match[0]
            sub_features = re.split(r'[|,;]\s*', match.strip())
            features.extend([f.strip() for f in sub_features if f.strip()])
    features = list(set(features))  # Remove duplicates
    logger.debug(f"Extracted features: {features}")
    return features

def read_file(uploaded_file):
    valid_extensions = {'.csv', '.xlsx', '.txt', '.pdf'}
    filename = secure_filename(uploaded_file.filename)
    extension = os.path.splitext(filename)[1].lower()
    
    if extension not in valid_extensions:
        logger.warning(f"Invalid file extension: {extension}")
        return {"error": f"Unsupported file type: {extension}. Supported types: {', '.join(valid_extensions)}"}
    
    try:
        contents = uploaded_file.read()
        logger.debug(f"Reading file: {filename}")
        
        if extension == '.xlsx':
            try:
                import openpyxl
                df = pd.read_excel(BytesIO(contents), engine='openpyxl')
                df = df.astype(object).where(pd.notnull(df), None)  # Convert to Python types
                return df
            except ImportError:
                logger.error("openpyxl not installed for .xlsx processing")
                return {"error": "Cannot process .xlsx files: openpyxl is not installed. Install it using 'pip install openpyxl'."}
        elif extension == '.csv':
            df = pd.read_csv(BytesIO(contents))
            df = df.astype(object).where(pd.notnull(df), None)
            return df
        elif extension == '.txt':
            try:
                text = contents.decode('utf-8')
                # First try to parse as delimited data
                for delimiter in [',', '\t', ';']:
                    try:
                        df = pd.read_csv(BytesIO(contents), delimiter=delimiter, on_bad_lines='skip')
                        if len(df.columns) > 1 and len(df) > 0:
                            df.columns = [f"col_{i}_{sanitize_table_name(str(col))}" for i, col in enumerate(df.columns)]
                            df = df.astype(object).where(pd.notnull(df), None)
                            return df
                    except:
                        continue
                
                # If no delimited format works, treat as plain text
                features = extract_features(text)
                if features:
                    df = pd.DataFrame({'feature': features})
                    return df
                else:
                    return {"error": "Could not extract structured data from text file. Please ensure it contains tabular data or clear feature lists."}
            except UnicodeDecodeError:
                return {"error": "Could not decode text file. Please ensure it's in UTF-8 format."}
        elif extension == '.pdf':
            try:
                text = ''
                with pdfplumber.open(BytesIO(contents)) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + '\n'
                
                if not text.strip():
                    # Fallback to PyPDF2
                    try:
                        pdf_reader = PyPDF2.PdfReader(BytesIO(contents))
                        for page in pdf_reader.pages:
                            text += page.extract_text() + '\n'
                    except:
                        pass
                
                if not text.strip():
                    return {"error": "Could not extract text from PDF. The file might be image-based or corrupted."}
                
                # Extract features from PDF text
                features = extract_features(text)
                if features:
                    df = pd.DataFrame({'feature': features})
                    df['source'] = 'PDF Document'
                    return df
                else:
                    return {"error": "Could not extract structured features from PDF. Please ensure it contains clear lists or tabular data."}
            except Exception as e:
                logger.error(f"Error processing PDF: {str(e)}")
                return {"error": f"Error processing PDF: {str(e)}"}
    except Exception as e:
        logger.error(f"Error reading file: {str(e)}")
        return {"error": f"Error reading file: {str(e)}"}

def create_table_in_postgres(df: pd.DataFrame, table_name: str):
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return False, connection

    cursor = connection.cursor()
    column_metadata = [(re.sub(r'[^a-zA-Z0-9_]', '_', col), "TEXT") for col in df.columns]
    
    for col_name, _ in column_metadata:
        if col_name.startswith('col_') or col_name == '':
            logger.warning(f"Generic or empty column name detected: {col_name}. Consider renaming columns for clarity.")

    column_definitions = ', '.join([f'"{name}" {dtype}' for name, dtype in column_metadata])
    
    create_table_query = f'CREATE TABLE IF NOT EXISTS "{table_name}" ({column_definitions});'

    try:
        cursor.execute(create_table_query)
        connection.commit()
        logger.debug(f"Table {table_name} created successfully")
        return True, column_metadata
    except Exception as e:
        connection.rollback()
        logger.error(f"Error creating table in PostgreSQL: {str(e)}")
        return False, {"error": f"Error creating table in PostgreSQL: {str(e)}"}
    finally:
        cursor.close()
        connection.close()

def insert_data_into_postgres(df: pd.DataFrame, table_name: str):
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return False, connection

    cursor = connection.cursor()

    try:
        for _, row in df.iterrows():
            values = tuple(str(val) if pd.notnull(val) else None for val in row)
            placeholders = ', '.join(['%s'] * len(values))
            insert_query = f'INSERT INTO "{table_name}" VALUES ({placeholders})'
            cursor.execute(insert_query, values)
        
        connection.commit()
        logger.debug(f"Data inserted into {table_name} successfully")
        return True, None
    except Exception as e:
        connection.rollback()
        logger.error(f"Error inserting data into PostgreSQL: {str(e)}")
        return False, {"error": f"Error inserting data into PostgreSQL: {str(e)}"}
    finally:
        cursor.close()
        connection.close()

def clean_sql_query(sql_query: str) -> str:
    """Clean and validate SQL query to prevent syntax errors"""
    # Remove extra whitespace and newlines
    sql_query = ' '.join(sql_query.split())
    
    # Remove trailing semicolons that might interfere with LIMIT
    sql_query = sql_query.rstrip(';')
    
    # Fix common syntax issues
    sql_query = re.sub(r'\s+', ' ', sql_query)  # Multiple spaces to single
    sql_query = re.sub(r'\(\s*;\s*\)', '()', sql_query)  # Remove semicolons in parentheses
    sql_query = re.sub(r';\s*LIMIT', ' LIMIT', sql_query)  # Fix semicolon before LIMIT
    
    return sql_query.strip()

def should_add_limit(sql_query: str) -> bool:
    """Determine if LIMIT clause should be added to the query"""
    sql_upper = sql_query.upper()
    
    # Don't add LIMIT if already present
    if 'LIMIT' in sql_upper:
        return False
    
    # Don't add LIMIT for aggregation queries
    if any(func in sql_upper for func in ['COUNT(', 'SUM(', 'AVG(', 'MAX(', 'MIN(']):
        return False
    
    # Don't add LIMIT for CREATE, INSERT, UPDATE, DELETE
    if any(cmd in sql_upper for cmd in ['CREATE', 'INSERT', 'UPDATE', 'DELETE', 'DROP']):
        return False
    
    # Don't add LIMIT if query ends with closing parenthesis (subquery)
    if sql_query.strip().endswith(')'):
        return False
    
    return True

def analyze_query_intent(question: str, columns: list, sample_text: str) -> dict:
    """Analyze user question to determine query intent and optimize SQL generation"""
    question_lower = question.lower()
    
    # Determine query type
    if any(word in question_lower for word in ['count', 'how many', 'total', 'number']):
        query_type = 'COUNT'
        description = 'Counting/aggregation query'
    elif any(word in question_lower for word in ['show', 'display', 'list', 'get', 'find', 'select']):
        query_type = 'SELECT'
        description = 'Data retrieval query'
    elif any(word in question_lower for word in ['skill', 'experience', 'education', 'qualification']):
        query_type = 'CONTENT_SEARCH'
        description = 'Resume/content analysis query'
    elif any(word in question_lower for word in ['all', 'everything', 'data', 'records']):
        query_type = 'FULL_TABLE'
        description = 'Complete data exploration'
    else:
        query_type = 'GENERAL'
        description = 'General information query'
    
    # Extract keywords
    keywords = []
    for word in question.split():
        if len(word) > 3 and word.lower() not in ['what', 'where', 'when', 'show', 'find', 'the', 'and', 'or']:
            keywords.append(word.lower())
    
    # Suggest best column
    best_column = 'full_text' if 'full_text' in columns else columns[0] if columns else None
    
    return {
        'type': query_type,
        'description': description,
        'keywords': keywords[:3],  # Top 3 keywords
        'best_column': best_column
    }

def generate_sql_query(user_question: str, table_name: str, column_metadata):
    if not user_question.strip() or len(user_question.strip()) < 3:
        return None, {"error": "Query is too short or empty. Please provide a specific question."}
    
    if user_question.strip().lower() in ['hii', 'hi', 'hello']:
        return None, {"error": "Query is too vague (e.g., 'hi'). Please ask a specific question about the data."}

    original_column_names = [name for name, _ in column_metadata]
    column_info = ", ".join([f"{name} ({dtype})" for name, dtype in column_metadata])
    
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return None, connection

    total_rows = 0
    sample_rows = []
    full_text_sample = ''
    try:
        cursor = connection.cursor()
        cursor.execute(f'SELECT COUNT(*) FROM "{table_name}"')
        total_rows = cursor.fetchone()[0]
        cursor.execute(f'SELECT * FROM "{table_name}" LIMIT 3')
        sample_rows = cursor.fetchall()
        sample_data = [
            {original_column_names[i]: str(val) if val is not None else 'NULL' for i, val in enumerate(row)}
            for row in sample_rows
        ]
        if 'full_text' in original_column_names:
            cursor.execute(f'SELECT "full_text" FROM "{table_name}" LIMIT 1')
            full_text_sample = cursor.fetchone()[0] or ''
            logger.debug(f"Full text sample (first 500 chars): {full_text_sample[:500]}")
        cursor.close()
    except Exception as e:
        connection.close()
        logger.error(f"Error fetching table info: {str(e)}")
        return None, {"error": f"Error fetching table info: {str(e)}"}
    finally:
        connection.close()
    
    if not full_text_sample and 'full_text' in original_column_names:
        logger.warning("Full text column is empty. Queries may not return expected results.")
    
    limit_clause = 'LIMIT 50' if total_rows > 50 else f'LIMIT {total_rows}'
    sample_data_str = "\n".join([f"Row {i+1}: {row}" for i, row in enumerate(sample_data)]) if sample_data else "No sample data available."
    
    features = extract_features(full_text_sample) if full_text_sample else []
    features_str = "\n".join([f"- {f}" for f in features]) if features else "No features identified."
    
    is_resume_query = any(keyword in user_question.lower() for keyword in ['skill', 'experience', 'education', 'qualification', 'proficiency'])
    is_feature_query = any(keyword in user_question.lower() for keyword in ['feature', 'functionality', 'capability'])
    # Analyze query intent for better optimization
    query_intent = analyze_query_intent(user_question, original_column_names, full_text_sample)
    
    # Get feedback patterns for RLHF
    feedback_patterns = get_feedback_patterns(user_question)
    
    prompt = (
        f"You are an expert PostgreSQL query generator specializing in natural language to SQL conversion. "
        f"Generate the most OPTIMIZED and ACCURATE SQL query for the given question.\n\n"
        f"DATABASE SCHEMA:\n"
        f"Table: '{table_name}' ({total_rows} rows)\n"
        f"Columns: {column_info}\n"
        f"Available columns: {', '.join(original_column_names)}\n\n"
        f"SAMPLE DATA (for context):\n{sample_data_str}\n\n"
        f"QUERY INTENT: {query_intent['type']} - {query_intent['description']}\n"
    )
    
    if feedback_patterns:
        prompt += f"\nLEARNED PATTERNS (from user feedback):\n{feedback_patterns}\n"
    
    if 'full_text' in original_column_names:
        prompt += f"DOCUMENT CONTENT PREVIEW:\n{full_text_sample[:500]}...\n\n"
        if is_resume_query or is_feature_query:
            prompt += f"EXTRACTED KEYWORDS: {features_str}\n\n"
    
    prompt += (
        f"OPTIMIZATION RULES:\n"
        f"1. COLUMN USAGE: Only use these exact columns: {', '.join(original_column_names)}\n"
        f"2. SYNTAX: Always enclose column names in double quotes\n"
        f"3. TEXT SEARCH: Use ILIKE with % wildcards for flexible matching\n"
        f"4. PERFORMANCE: Add {limit_clause} unless using aggregation (COUNT, SUM, AVG)\n"
        f"5. DATA TYPES: All columns are TEXT - use CAST() for numeric operations\n"
        f"6. NULL HANDLING: Use IS NULL/IS NOT NULL for null checks\n\n"
        f"QUERY PATTERNS:\n"
        f"• Data Exploration: SELECT * FROM \"{table_name}\" {limit_clause}\n"
        f"• Counting: SELECT COUNT(*) FROM \"{table_name}\" WHERE condition\n"
        f"• Text Search: SELECT * FROM \"{table_name}\" WHERE \"column\" ILIKE '%keyword%'\n"
        f"• Content Analysis: SELECT \"full_text\" FROM \"{table_name}\" WHERE \"full_text\" ILIKE '%{query_intent['keywords'][0] if query_intent['keywords'] else 'content'}%'\n"
        f"• Numeric Filter: SELECT * FROM \"{table_name}\" WHERE CAST(\"column\" AS NUMERIC) > value\n\n"
        f"USER QUESTION: '{user_question}'\n\n"
        f"GENERATE OPTIMIZED SQL (query only, no explanations):\n"
    )

    try:
        generated_response = gemini.run(parts=[prompt])
        logger.debug(f"Generated SQL query: {generated_response}")
        
        if "replies" in generated_response and generated_response["replies"]:
            sql_query = generated_response["replies"][0].strip()
            
            sql_query = re.sub(r'^```sql\s*|\s*```$', '', sql_query, flags=re.MULTILINE).strip()
            
            if not sql_query:
                return None, {"error": "Generated SQL query is empty."}
            
            if not sql_query.upper().startswith(('SELECT', 'INSERT', 'UPDATE', 'DELETE')):
                return None, {"error": f"Invalid SQL query generated: {sql_query}"}
            
            for column in original_column_names:
                sql_query = re.sub(
                    r'(?<!")' + re.escape(column) + r'(?!")', 
                    f'"{column}"', 
                    sql_query
                )
            
            # Clean and validate SQL query
            sql_query = clean_sql_query(sql_query)
            
            # Add LIMIT clause only for valid SELECT statements
            if (should_add_limit(sql_query) and 
                sql_query.upper().startswith('SELECT')):
                sql_query += f' {limit_clause}'
            
            logger.debug(f"Final SQL query: {sql_query}")
            return sql_query, None
        
        return None, {"error": "Failed to generate SQL query: No response from generator."}
    
    except Exception as e:
        logger.error(f"Error generating SQL query: {str(e)}")
        return None, {"error": f"Error generating SQL query: {str(e)}"}

def execute_sql_query(sql_query: str):
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return None, connection

    try:
        cursor = connection.cursor()
        cursor.execute(sql_query)
        
        columns = [desc[0] for desc in cursor.description]
        rows = cursor.fetchall()
        
        result = [dict(zip(columns, row)) for row in rows]
        logger.debug(f"SQL query executed successfully: {sql_query}")
        return result, None
    
    except Exception as e:
        logger.error(f"Error executing SQL query: {str(e)}")
        return None, {"error": f"Error executing SQL query: {str(e)}"}
    finally:
        cursor.close()
        connection.close()

def generate_conversational_response(user_question: str, query_results: list):
    results_str = "\n".join([str(row) for row in query_results])
    is_resume_query = any(keyword in user_question.lower() for keyword in ['skill', 'experience', 'education', 'qualification', 'proficiency'])
    is_feature_query = any(keyword in user_question.lower() for keyword in ['feature', 'functionality', 'capability'])
    
    if (is_resume_query or is_feature_query) and query_results and 'full_text' in query_results[0]:
        full_text = query_results[0]['full_text']
        features = extract_features(full_text)
        if features:
            response = f"Your {'skills' if is_resume_query else 'features'} include:\n" + "\n".join([f"- {f}" for f in features])
        else:
            response = f"No {'skills' if is_resume_query else 'features'} were identified. Try rephrasing, e.g., 'List my {'experience' if is_resume_query else 'capabilities'}'."
    else:
        response = f"Based on your question '{user_question}', here are the results: {results_str}"
    
    prompt = (
        f"The user asked: '{user_question}'. "
        f"Here are the query results: {results_str}. "
        f"Initial response: {response}. "
        f"Refine the response to be concise, conversational, and clear. "
        f"For resume or feature-related questions, list items in bullet points. "
        f"If no items are found, suggest rephrasing."
    )

    try:
        generated_response = gemini.run(parts=[prompt])
        logger.debug(f"Generated conversational response: {generated_response}")
        
        if "replies" in generated_response and generated_response["replies"]:
            return generated_response["replies"][0], None
        
        return response, None
    
    except Exception as e:
        logger.error(f"Error generating conversational response: {str(e)}")
        return response, None

@app.route('/')
def serve_frontend():
    try:
        logger.debug(f"Serving index.html from {app.static_folder}")
        return send_from_directory(app.static_folder, 'index.html')
    except Exception as e:
        logger.error(f"Error serving index.html: {str(e)}")
        return jsonify({"error": f"Failed to serve frontend: {str(e)}"}), 500

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        logger.warning("No file provided in upload request")
        return jsonify({"detail": "No file provided."}), 400
    
    file = request.files['file']
    if not file.filename:
        logger.warning("No file selected in upload request")
        return jsonify({"detail": "No file selected."}), 400

    file.seek(0, os.SEEK_END)
    file_size = file.tell()
    file.seek(0)

    df = read_file(file)
    if isinstance(df, dict):
        logger.error(f"Upload failed: {df['error']}")
        return jsonify({"detail": df["error"]}), 500
    
    # Calculate text length properly for different file types
    full_text_length = 0
    if not df.empty:
        if 'full_text' in df.columns:
            # For PDF and text files with full_text column
            full_text_series = df['full_text'].dropna()
            if not full_text_series.empty:
                full_text_length = int(full_text_series.str.len().sum())
        else:
            # For CSV and other structured files, calculate total text length
            text_columns = df.select_dtypes(include=['object']).fillna('')
            if not text_columns.empty:
                full_text_length = int(text_columns.astype(str).apply(lambda x: x.str.len().sum()).sum())
    
    if df is None or df.empty:
        logger.warning("Unsupported file format or empty file uploaded")
        return jsonify({"detail": "Unsupported file format or empty file. Ensure the file contains readable data."}), 400
    
    if full_text_length < 50 and 'full_text' in df.columns:
        logger.warning(f"Minimal text extracted (length: {full_text_length}). Queries may be limited.")
    
    df = df.fillna(pd.NA).replace({pd.NA: None})

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    table_name = sanitize_table_name(f"{file.filename.split('.')[0]}_{timestamp}")
    
    table_created, column_metadata = create_table_in_postgres(df, table_name)
    if not table_created:
        logger.error(f"Table creation failed: {column_metadata['error']}")
        return jsonify({"detail": column_metadata["error"]}), 500
    
    inserted, error = insert_data_into_postgres(df, table_name)
    if not inserted:
        logger.error(f"Data insertion failed: {error['error']}")
        return jsonify({"detail": error["error"]}), 500
    
    table_metadata["tables"][table_name] = column_metadata
    table_metadata["current_table"] = table_name
    
    metadata = {
        "file_size": f"{file_size / 1024:.2f} KB",
        "row_count": int(len(df)),
        "column_count": int(len(df.columns)),
        "text_length": int(full_text_length)
    }

    logger.debug(f"File uploaded and processed successfully. Current table set to: {table_name}")
    return jsonify({
        "data": df.head().to_dict(orient="records"),
        "table_name": table_name,
        "metadata": metadata
    })

def get_existing_tables_from_db():
    """Get all existing tables from the database"""
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return []
    
    cursor = connection.cursor()
    try:
        # Get all tables except system tables
        cursor.execute("""
            SELECT table_name 
            FROM information_schema.tables 
            WHERE table_schema = 'public' 
            AND table_name NOT IN ('query_history', 'feedback_patterns')
            ORDER BY table_name;
        """)
        tables = cursor.fetchall()
        
        table_list = []
        for (table_name,) in tables:
            # Get column information for each table
            cursor.execute("""
                SELECT column_name 
                FROM information_schema.columns 
                WHERE table_name = %s 
                AND table_schema = 'public'
                ORDER BY ordinal_position;
            """, (table_name,))
            columns = [col[0] for col in cursor.fetchall()]
            table_list.append({"name": table_name, "columns": columns})
            
            # Update table_metadata if not already present
            if table_name not in table_metadata["tables"]:
                table_metadata["tables"][table_name] = [(col, "TEXT") for col in columns]
        
        return table_list
    except Exception as e:
        logger.error(f"Error fetching existing tables: {str(e)}")
        return []
    finally:
        cursor.close()
        connection.close()

@app.route('/tables', methods=['GET'])
def get_tables():
    try:
        # Get tables from memory and database
        memory_tables = [{"name": name, "columns": [col[0] for col in metadata]} for name, metadata in table_metadata["tables"].items()]
        db_tables = get_existing_tables_from_db()
        
        # Merge and deduplicate
        all_tables = {}
        for table in memory_tables + db_tables:
            all_tables[table["name"]] = table
        
        tables = list(all_tables.values())
        logger.debug(f"Fetched {len(tables)} tables successfully")
        return jsonify({"tables": tables})
    except Exception as e:
        logger.error(f"Error fetching tables: {str(e)}")
        return jsonify({"detail": f"Error fetching tables: {str(e)}"}), 500

@app.route('/delete_query', methods=['POST'])
def delete_query():
    data = request.get_json()
    if not data or 'query_id' not in data:
        return jsonify({"detail": "No query ID provided."}), 400
    
    query_id = data['query_id']
    
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return jsonify({"detail": connection["error"]}), 500
    
    cursor = connection.cursor()
    try:
        cursor.execute('DELETE FROM query_history WHERE id = %s', (query_id,))
        connection.commit()
        
        logger.debug(f"Query {query_id} deleted successfully")
        return jsonify({"message": f"Query deleted successfully"})
        
    except Exception as e:
        connection.rollback()
        logger.error(f"Error deleting query: {str(e)}")
        return jsonify({"detail": f"Error deleting query: {str(e)}"}), 500
    finally:
        cursor.close()
        connection.close()

@app.route('/delete_table', methods=['POST'])
def delete_table():
    data = request.get_json()
    if not data or 'table_name' not in data:
        return jsonify({"detail": "No table name provided."}), 400
    
    table_name = data['table_name']
    
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return jsonify({"detail": connection["error"]}), 500
    
    cursor = connection.cursor()
    try:
        # Delete the table from database
        cursor.execute(f'DROP TABLE IF EXISTS "{table_name}"')
        connection.commit()
        
        # Remove from memory
        if table_name in table_metadata["tables"]:
            del table_metadata["tables"][table_name]
        
        # If this was the current table, clear it
        if table_metadata["current_table"] == table_name:
            table_metadata["current_table"] = None
        
        logger.debug(f"Table {table_name} deleted successfully")
        return jsonify({"message": f"Table {table_name} deleted successfully"})
        
    except Exception as e:
        connection.rollback()
        logger.error(f"Error deleting table: {str(e)}")
        return jsonify({"detail": f"Error deleting table: {str(e)}"}), 500
    finally:
        cursor.close()
        connection.close()

@app.route('/select_table', methods=['POST'])
def select_table():
    data = request.get_json()
    if not data or 'table_name' not in data:
        logger.warning("No table name provided")
        return jsonify({"detail": "No table name provided."}), 400
    
    table_name = data['table_name']
    
    # If table not in memory, try to load it from database
    if table_name not in table_metadata["tables"]:
        connection = connect_to_postgres()
        if not isinstance(connection, dict):
            cursor = connection.cursor()
            try:
                # Check if table exists in database
                cursor.execute("""
                    SELECT column_name 
                    FROM information_schema.columns 
                    WHERE table_name = %s 
                    AND table_schema = 'public'
                    ORDER BY ordinal_position;
                """, (table_name,))
                columns = cursor.fetchall()
                
                if columns:
                    # Load table metadata
                    table_metadata["tables"][table_name] = [(col[0], "TEXT") for col in columns]
                    logger.debug(f"Loaded table metadata for {table_name} from database")
                else:
                    logger.warning(f"Table {table_name} not found in database")
                    return jsonify({"detail": "Invalid table name."}), 400
            except Exception as e:
                logger.error(f"Error loading table metadata: {str(e)}")
                return jsonify({"detail": f"Error loading table: {str(e)}"}), 500
            finally:
                cursor.close()
                connection.close()
        else:
            return jsonify({"detail": "Database connection failed."}), 500
    
    table_metadata["current_table"] = table_name
    logger.debug(f"Selected table: {table_name}")
    return jsonify({"message": f"Selected table: {table_name}"})

@app.route('/query', methods=['POST'])
def process_query():
    data = request.get_json()
    if not data or 'question' not in data:
        logger.warning("No question provided in query request")
        return jsonify({"detail": "No question provided."}), 400
    
    if not table_metadata["current_table"]:
        logger.warning("No table selected for query")
        return jsonify({"detail": "No table selected. Please upload or select a file."}), 400
    
    sql_query, error = generate_sql_query(
        data['question'],
        table_metadata["current_table"],
        table_metadata["tables"][table_metadata["current_table"]]
    )
    if sql_query is None:
        logger.error(f"SQL query generation failed: {error['error']}")
        return jsonify({"detail": error["error"]}), 500
    
    results, error = execute_sql_query(sql_query)
    if results is None:
        logger.error(f"SQL query execution failed: {error['error']}")
        return jsonify({"detail": error["error"]}), 500
    
    conversational_response, error = generate_conversational_response(data['question'], results)
    if conversational_response is None:
        logger.error(f"Conversational response generation failed: {error['error']}")
        return jsonify({"detail": error["error"]}), 500
    
    connection = connect_to_postgres()
    query_id = None
    if not isinstance(connection, dict):
        cursor = connection.cursor()
        insert_query = """
        INSERT INTO query_history (table_name, question, sql_query, results)
        VALUES (%s, %s, %s, %s) RETURNING id;
        """
        try:
            cursor.execute(insert_query, (table_metadata["current_table"], data['question'], sql_query, json.dumps(results)))
            query_id = cursor.fetchone()[0]
            connection.commit()
            logger.debug(f"Query stored in history with ID: {query_id}")
        except Exception as e:
            connection.rollback()
            logger.error(f"Error storing query in history: {str(e)}")
        finally:
            cursor.close()
            connection.close()

    logger.debug("Query processed successfully")
    return jsonify({
        "sql_query": sql_query,
        "results": results,
        "conversational_response": conversational_response,
        "query_id": query_id
    })

@app.route('/query_history', methods=['GET'])
def get_query_history():
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        logger.error(f"Failed to connect to PostgreSQL: {connection['error']}")
        return jsonify({"detail": connection["error"]}), 500
    
    cursor = connection.cursor()
    select_query = """
    SELECT id, table_name, question, sql_query, results, timestamp, feedback
    FROM query_history
    ORDER BY timestamp DESC
    LIMIT 50;
    """
    try:
        cursor.execute(select_query)
        columns = [desc[0] for desc in cursor.description]
        rows = cursor.fetchall()
        history = [dict(zip(columns, row)) for row in rows]
        logger.debug("Fetched query history successfully")
        return jsonify({"history": history})
    except Exception as e:
        logger.error(f"Error fetching query history: {str(e)}")
        return jsonify({"detail": f"Error fetching query history: {str(e)}"}), 500
    finally:
        cursor.close()
        connection.close()

@app.route('/feedback', methods=['POST'])
def submit_feedback():
    data = request.get_json()
    if not data or 'query_id' not in data or 'feedback' not in data:
        return jsonify({"detail": "Query ID and feedback are required."}), 400
    
    query_id = data['query_id']
    feedback = data['feedback']  # 1 for like, -1 for dislike
    
    if feedback not in [1, -1]:
        return jsonify({"detail": "Feedback must be 1 (like) or -1 (dislike)."}), 400
    
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return jsonify({"detail": connection["error"]}), 500
    
    cursor = connection.cursor()
    try:
        # Update feedback in query_history
        cursor.execute("""
            UPDATE query_history 
            SET feedback = %s, feedback_timestamp = CURRENT_TIMESTAMP 
            WHERE id = %s
        """, (feedback, query_id))
        
        # Get the query details for pattern learning
        cursor.execute("""
            SELECT question, sql_query 
            FROM query_history 
            WHERE id = %s
        """, (query_id,))
        
        result = cursor.fetchone()
        if result:
            question, sql_query = result
            
            # Extract question pattern (keywords)
            question_pattern = ' '.join([word.lower() for word in question.split() 
                                       if len(word) > 3 and word.lower() not in 
                                       ['what', 'show', 'find', 'the', 'and', 'or', 'with']])
            
            # Update or insert feedback pattern
            if feedback == 1:  # Positive feedback
                cursor.execute("""
                    SELECT id, feedback_score, usage_count FROM feedback_patterns 
                    WHERE question_pattern = %s
                """, (question_pattern,))
                existing = cursor.fetchone()
                
                if existing:
                    pattern_id, current_score, current_count = existing
                    new_score = (current_score * current_count + 1.0) / (current_count + 1)
                    cursor.execute("""
                        UPDATE feedback_patterns 
                        SET feedback_score = %s, usage_count = %s, successful_query = %s, updated_at = CURRENT_TIMESTAMP
                        WHERE id = %s
                    """, (new_score, current_count + 1, sql_query, pattern_id))
                else:
                    cursor.execute("""
                        INSERT INTO feedback_patterns (question_pattern, successful_query, feedback_score, usage_count)
                        VALUES (%s, %s, 1.0, 1)
                    """, (question_pattern, sql_query))
            else:  # Negative feedback
                cursor.execute("""
                    SELECT id, feedback_score, usage_count FROM feedback_patterns 
                    WHERE question_pattern = %s
                """, (question_pattern,))
                existing = cursor.fetchone()
                
                if existing:
                    pattern_id, current_score, current_count = existing
                    new_score = (current_score * current_count - 0.5) / (current_count + 1)
                    cursor.execute("""
                        UPDATE feedback_patterns 
                        SET feedback_score = %s, usage_count = %s, updated_at = CURRENT_TIMESTAMP
                        WHERE id = %s
                    """, (new_score, current_count + 1, pattern_id))
                else:
                    cursor.execute("""
                        INSERT INTO feedback_patterns (question_pattern, successful_query, feedback_score, usage_count)
                        VALUES (%s, %s, -0.5, 1)
                    """, (question_pattern, sql_query))
        
        connection.commit()
        logger.debug(f"Feedback {feedback} recorded for query {query_id}")
        return jsonify({"message": "Feedback recorded successfully"})
        
    except Exception as e:
        connection.rollback()
        logger.error(f"Error recording feedback: {str(e)}")
        return jsonify({"detail": f"Error recording feedback: {str(e)}"}), 500
    finally:
        cursor.close()
        connection.close()

def get_feedback_patterns(question: str) -> str:
    """Get successful query patterns based on previous feedback"""
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return ""
    
    cursor = connection.cursor()
    try:
        # Check if feedback_patterns table exists
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables 
                WHERE table_schema = 'public' 
                AND table_name = 'feedback_patterns'
            )
        """)
        table_exists = cursor.fetchone()[0]
        
        if not table_exists:
            return ""
        
        # Extract keywords from current question
        keywords = [word.lower() for word in question.split() 
                   if len(word) > 3 and word.lower() not in 
                   ['what', 'show', 'find', 'the', 'and', 'or', 'with']]
        
        if not keywords:
            return ""
        
        # Find similar patterns with positive feedback
        pattern_conditions = []
        for keyword in keywords:
            pattern_conditions.append(f"question_pattern ILIKE '%{keyword}%'")
        
        if pattern_conditions:
            cursor.execute(f"""
                SELECT successful_query, feedback_score, usage_count
                FROM feedback_patterns
                WHERE feedback_score > 0.3
                AND ({' OR '.join(pattern_conditions)})
                ORDER BY feedback_score DESC, usage_count DESC
                LIMIT 3
            """)
        
        patterns = cursor.fetchall()
        
        if patterns:
            successful_queries = []
            for query, score, count in patterns:
                successful_queries.append(f"Similar successful query (score: {score:.2f}): {query}")
            
            return "\n".join(successful_queries)
        
        return ""
        
    except Exception as e:
        logger.error(f"Error fetching feedback patterns: {str(e)}")
        return ""
    finally:
        cursor.close()
        connection.close()

@app.route('/enhanced_feedback', methods=['POST'])
def submit_enhanced_feedback():
    data = request.get_json()
    if not data or 'query_id' not in data or 'rating' not in data:
        return jsonify({"detail": "Query ID and rating are required."}), 400
    
    query_id = data['query_id']
    rating = data['rating']  # 1-5 stars
    categories = data.get('categories', [])
    comment = data.get('comment', '')
    
    if rating not in [1, 2, 3, 4, 5]:
        return jsonify({"detail": "Rating must be between 1 and 5."}), 400
    
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return jsonify({"detail": connection["error"]}), 500
    
    cursor = connection.cursor()
    try:
        # Update enhanced feedback in query_history
        cursor.execute("""
            UPDATE query_history 
            SET rating = %s, categories = %s, comment = %s, feedback_timestamp = CURRENT_TIMESTAMP 
            WHERE id = %s
        """, (rating, categories, comment, query_id))
        
        # Get the query details for pattern learning
        cursor.execute("""
            SELECT question, sql_query 
            FROM query_history 
            WHERE id = %s
        """, (query_id,))
        
        result = cursor.fetchone()
        if result:
            question, sql_query = result
            
            # Extract question pattern (keywords)
            question_pattern = ' '.join([word.lower() for word in question.split() 
                                       if len(word) > 3 and word.lower() not in 
                                       ['what', 'show', 'find', 'the', 'and', 'or', 'with']])
            
            # Convert rating to feedback score (1-2 = negative, 3 = neutral, 4-5 = positive)
            feedback_score = (rating - 3) / 2.0  # Maps 1->-1, 2->-0.5, 3->0, 4->0.5, 5->1
            
            # Update or insert feedback pattern
            cursor.execute("""
                SELECT id, feedback_score, usage_count FROM feedback_patterns 
                WHERE question_pattern = %s
            """, (question_pattern,))
            existing = cursor.fetchone()
            
            if existing:
                pattern_id, current_score, current_count = existing
                new_score = (current_score * current_count + feedback_score) / (current_count + 1)
                cursor.execute("""
                    UPDATE feedback_patterns 
                    SET feedback_score = %s, usage_count = %s, successful_query = %s, updated_at = CURRENT_TIMESTAMP
                    WHERE id = %s
                """, (new_score, current_count + 1, sql_query if rating >= 4 else None, pattern_id))
            else:
                cursor.execute("""
                    INSERT INTO feedback_patterns (question_pattern, successful_query, feedback_score, usage_count)
                    VALUES (%s, %s, %s, 1)
                """, (question_pattern, sql_query if rating >= 4 else None, feedback_score))
        
        connection.commit()
        logger.debug(f"Enhanced feedback recorded for query {query_id}: rating={rating}, categories={categories}")
        return jsonify({"message": "Enhanced feedback recorded successfully"})
        
    except Exception as e:
        connection.rollback()
        logger.error(f"Error recording enhanced feedback: {str(e)}")
        return jsonify({"detail": f"Error recording enhanced feedback: {str(e)}"}), 500
    finally:
        cursor.close()
        connection.close()

@app.route('/feedback_analytics', methods=['GET'])
def get_feedback_analytics():
    connection = connect_to_postgres()
    if isinstance(connection, dict):
        return jsonify({"detail": connection["error"]}), 500
    
    cursor = connection.cursor()
    try:
        # Get total feedback count
        cursor.execute("SELECT COUNT(*) FROM query_history WHERE rating IS NOT NULL")
        total_feedback = cursor.fetchone()[0] or 0
        
        if total_feedback == 0:
            # Return default values when no feedback exists
            return jsonify({
                "total_feedback": 0,
                "avg_rating": 0,
                "satisfaction_rate": 0,
                "common_issues": []
            })
        
        # Get average rating
        cursor.execute("SELECT AVG(rating) FROM query_history WHERE rating IS NOT NULL")
        avg_rating_result = cursor.fetchone()[0]
        avg_rating = round(float(avg_rating_result), 1) if avg_rating_result else 0
        
        # Get satisfaction rate (4-5 star ratings)
        cursor.execute("""
            SELECT 
                CASE 
                    WHEN COUNT(*) = 0 THEN 0
                    ELSE COUNT(CASE WHEN rating >= 4 THEN 1 END) * 100.0 / COUNT(*)
                END as satisfaction_rate
            FROM query_history 
            WHERE rating IS NOT NULL
        """)
        satisfaction_result = cursor.fetchone()[0]
        satisfaction_rate = round(float(satisfaction_result), 1) if satisfaction_result else 0
        
        # Get most common issues (from categories)
        cursor.execute("""
            SELECT unnest(categories) as category, COUNT(*) as count
            FROM query_history 
            WHERE categories IS NOT NULL AND array_length(categories, 1) > 0
            GROUP BY category
            ORDER BY count DESC
            LIMIT 5
        """)
        common_issues = cursor.fetchall()
        
        analytics = {
            "total_feedback": total_feedback,
            "avg_rating": avg_rating,
            "satisfaction_rate": satisfaction_rate,
            "common_issues": [{
                "category": issue[0],
                "count": issue[1]
            } for issue in common_issues]
        }
        
        logger.debug(f"Feedback analytics: {analytics}")
        return jsonify(analytics)
        
    except Exception as e:
        logger.error(f"Error fetching feedback analytics: {str(e)}")
        return jsonify({
            "total_feedback": 0,
            "avg_rating": 0,
            "satisfaction_rate": 0,
            "common_issues": []
        })
    finally:
        cursor.close()
        connection.close()

if __name__ == '__main__':
    logger.info("Starting HumanSQL - Natural Language to SQL Converter")
    app.run(host='0.0.0.0', port=8080, debug=True)